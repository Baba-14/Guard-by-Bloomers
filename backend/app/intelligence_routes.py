"""Consent-aware collection, review and dataset administration routes."""

import csv
import hashlib
import hmac
import io
import json
from datetime import datetime, timedelta, timezone
from typing import Annotated, Literal
from uuid import UUID

from fastapi import APIRouter, Depends, File, HTTPException, Query, Request, UploadFile, status
from pydantic import BaseModel, Field
from sqlalchemy import func, select
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session, selectinload

from .auth import get_optional_user, require_analyst
from .config import get_settings
from .database import get_db
from .analysis.rules import analyse_text
from .intelligence import PHONE_PATTERN, URL_PATTERN, deidentify, normalize_phone, normalize_url
from .models import (
    AuditLog, Brand, DatasetItem, DatasetLabel, DatasetSource, DatasetSourceType, DatasetVersion,
    DatasetVersionItem, Feedback, FraudCategory, FraudReport, PhoneIntelligence, ReportStatus,
    ReviewQueueItem, ReviewStatus, UrlIntelligence, User, UserRole,
)

router = APIRouter(prefix="/v1", tags=["Guard Intelligence"])
Db = Annotated[Session, Depends(get_db)]
Analyst = Annotated[User, Depends(require_analyst)]
OptionalUser = Annotated[User | None, Depends(get_optional_user)]


class ReportRequest(BaseModel):
    report_type: str = Field(min_length=2, max_length=100)
    category: str = Field(min_length=2, max_length=100)
    identifier: str | None = Field(default=None, max_length=500)
    description: str = Field(min_length=10, max_length=20000)
    contribution_consent: Literal[True]
    consent_version: str = Field(default="guard-report-v1", max_length=50)
    country_code: str = Field(default="GH", min_length=2, max_length=2)
    channel: str | None = Field(default=None, max_length=50)


class ManualDatasetRequest(BaseModel):
    content: str = Field(min_length=2, max_length=20000)
    label: DatasetLabel = DatasetLabel.uncertain
    source_id: UUID
    category_id: UUID | None = None
    channel: str | None = None
    country_code: str | None = "GH"
    is_ghana_specific: bool = True
    language: str | None = "en"
    notes: str | None = None


class ReviewDecisionRequest(BaseModel):
    decision: DatasetLabel
    category_id: UUID | None = None
    notes: str | None = Field(default=None, max_length=4000)
    approve: bool = True


class FeedbackRequest(BaseModel):
    helpful: bool
    kind: str = Field(min_length=2, max_length=30)
    result_level: str = Field(min_length=2, max_length=40)


class AdminUserUpdateRequest(BaseModel):
    role: UserRole | None = None
    is_active: bool | None = None


class TrainingDatasetRequest(BaseModel):
    notes: str | None = Field(default=None, max_length=2000)


REPORT_CATEGORY_DESCRIPTIONS = {
    "Mobile Money Fraud": "Fraud involving mobile money accounts, transfers or reversals",
    "Fake Payment Alert or Screenshot": "Fabricated or manipulated payment confirmation",
    "Wrong Transfer or Reversal Request": "Requests to return or reverse a supposedly mistaken transfer",
    "WhatsApp Account Takeover": "Compromised WhatsApp accounts used to request money or information",
    "WhatsApp Impersonation": "WhatsApp accounts pretending to be a trusted person or organisation",
    "OTP or Verification Code Request": "Attempts to obtain one-time passwords or verification codes",
    "Phishing": "Credential or payment capture through deceptive links or messages",
    "Bank Impersonation": "People or messages pretending to represent a bank",
    "Fake Online Shop": "Fraudulent shops, merchants or product listings",
    "Investment Fraud": "False or misleading investment offers",
    "Other": "Reported fraud that does not yet match a specific category",
}


def item_hash(content: str) -> str:
    return hashlib.sha256(" ".join(content.lower().split()).encode()).hexdigest()


def submission_fingerprint(host: str, at: datetime | None = None) -> str:
    """Return a rotating one-way identifier; the raw address is never stored."""
    day = (at or datetime.now(timezone.utc)).date().isoformat()
    secret = get_settings().require_jwt_secret().encode()
    return hmac.new(secret, f"{day}:{host}".encode(), hashlib.sha256).hexdigest()


def category_for(db: Session, value: str) -> FraudCategory | None:
    normalized = value.lower().strip().replace(" ", "-")
    category = db.scalar(select(FraudCategory).where((FraudCategory.slug == normalized) | (func.lower(FraudCategory.name) == value.lower().strip())))
    canonical_name = next((name for name in REPORT_CATEGORY_DESCRIPTIONS if name.lower() == value.lower().strip()), None)
    if category is None and canonical_name is not None:
        category = FraudCategory(
            name=canonical_name,
            slug=canonical_name.lower().replace(" ", "-"),
            description=REPORT_CATEGORY_DESCRIPTIONS[canonical_name],
        )
        db.add(category)
        db.flush()
    return category


def serialize_item(item: DatasetItem) -> dict:
    return {
        "id": str(item.id), "content": item.content, "content_type": item.content_type,
        "channel": item.channel, "label": item.label.value, "category": item.category.name if item.category else None,
        "category_id": str(item.category_id) if item.category_id else None,
        "source": item.source.name, "source_type": item.source.source_type.value,
        "country_code": item.country_code, "is_ghana_specific": item.is_ghana_specific,
        "verification_status": item.verification_status.value, "privacy_status": item.privacy_status,
        "created_at": item.created_at,
    }


def serialize_user(user: User) -> dict:
    return {
        "id": str(user.id),
        "email": user.email,
        "full_name": user.profile.full_name,
        "role": user.profile.role.value,
        "is_active": user.is_active,
        "email_verified": user.email_verified,
        "created_at": user.created_at,
    }


def create_candidate(
    db: Session, *, content: str, source: DatasetSource, label: DatasetLabel = DatasetLabel.uncertain,
    category_id: UUID | None = None, report_id: UUID | None = None, channel: str | None = None,
    country_code: str | None = "GH", is_ghana_specific: bool = True, consent: bool = False,
    metadata: dict | None = None,
) -> DatasetItem:
    clean_content, had_pii = deidentify(content)
    extracted_signals = [signal.key for signal in analyse_text(clean_content)]
    item_metadata = dict(metadata or {})
    item_metadata.setdefault("extracted_signals", extracted_signals)
    item = DatasetItem(
        content=clean_content, content_hash=item_hash(clean_content), channel=channel, label=label,
        category_id=category_id, source_id=source.id, fraud_report_id=report_id,
        country_code=country_code, is_ghana_specific=is_ghana_specific,
        verification_status=ReviewStatus.pending, contribution_consent=consent,
        privacy_status="deidentified" if had_pii else "reviewed_no_direct_pii",
        metadata_json=item_metadata,
    )
    db.add(item)
    db.flush()
    db.add(ReviewQueueItem(dataset_item_id=item.id, status=ReviewStatus.pending))
    return item


def sync_verified_entity(db: Session, item: DatasetItem, report: FraudReport) -> None:
    """Promote an approved report identifier into moderated reputation intelligence."""
    if not report.entity_value:
        return
    report_type = (report.entity_type or "").lower()
    raw_value = report.entity_value.strip()
    if "link" in report_type or "website" in report_type:
        candidate = URL_PATTERN.search(raw_value)
        if candidate is None:
            candidate = URL_PATTERN.search(report.description)
        if candidate is None:
            return
        normalized = normalize_url(candidate.group(0))
        known = db.scalar(select(UrlIntelligence).where(UrlIntelligence.normalized_url == normalized))
        if known is None:
            known = UrlIntelligence(
                url=candidate.group(0), normalized_url=normalized, source_id=item.source_id,
                verdict=DatasetLabel.fraud, risk_score=85, report_count=1, verified=True,
                metadata_json={"origin": "approved_user_report"},
            )
            db.add(known)
        else:
            known.verdict = DatasetLabel.fraud
            known.risk_score = max(known.risk_score or 0, 85)
            known.report_count = (known.report_count or 0) + 1
            known.verified = True
        return
    if "phone" in report_type or "whatsapp" in report_type or "number" in report_type:
        candidate = PHONE_PATTERN.search(raw_value) or PHONE_PATTERN.search(report.description)
        if candidate is None:
            return
        normalized = normalize_phone(candidate.group(0))
        known = db.scalar(select(PhoneIntelligence).where(PhoneIntelligence.sender_id == normalized))
        if known is None:
            known = PhoneIntelligence(
                sender_id=normalized, source_id=item.source_id, verdict=DatasetLabel.fraud,
                risk_score=85, report_count=1, verified=True,
                metadata_json={"origin": "approved_user_report"},
            )
            db.add(known)
        else:
            known.verdict = DatasetLabel.fraud
            known.risk_score = max(known.risk_score or 0, 85)
            known.report_count = (known.report_count or 0) + 1
            known.verified = True


@router.post("/contributions", status_code=status.HTTP_201_CREATED)
def submit_report(payload: ReportRequest, http_request: Request, db: Db, current_user: OptionalUser) -> dict:
    """Explicit reports contribute a de-identified candidate; ordinary checks never call this."""
    fingerprint = submission_fingerprint(http_request.client.host if http_request.client else "unknown")
    recent_count = db.scalar(
        select(func.count()).select_from(DatasetItem).where(
            DatasetItem.created_at >= datetime.now(timezone.utc) - timedelta(hours=1),
            DatasetItem.metadata_json["submission_fingerprint"].astext == fingerprint,
        )
    ) or 0
    if recent_count >= 5:
        raise HTTPException(status_code=429, detail="Too many reports from this connection. Please try again later.")

    category = category_for(db, payload.category)
    source = db.scalar(select(DatasetSource).where(DatasetSource.source_type == DatasetSourceType.user_report))
    if source is None:
        source = DatasetSource(name="Guard user reports", source_type=DatasetSourceType.user_report, description="Explicit user contributions")
        db.add(source)
        db.flush()
    report = FraudReport(
        reporter_id=current_user.id if current_user else None,
        category_id=category.id if category else None, description=payload.description,
        entity_type=payload.report_type, entity_value=payload.identifier,
        explicit_contribution_consent=True, consent_version=payload.consent_version,
        privacy_status="restricted_raw",
    )
    db.add(report)
    db.flush()
    try:
        create_candidate(
            db, content=payload.description, source=source,
            category_id=category.id if category else None, report_id=report.id,
            channel=payload.channel, country_code=payload.country_code.upper(),
            is_ghana_specific=payload.country_code.upper() == "GH", consent=True,
            metadata={"report_type": payload.report_type, "identifier_provided": bool(payload.identifier), "submission_fingerprint": fingerprint},
        )
        db.commit()
    except IntegrityError:
        db.rollback()
        raise HTTPException(status_code=409, detail="This report content is already awaiting review")
    return {"id": str(report.id), "reference": f"GRD-{str(report.id)[:8].upper()}", "status": "pending_review"}


@router.post("/feedback", status_code=status.HTTP_201_CREATED)
def submit_feedback(payload: FeedbackRequest, db: Db) -> dict:
    """Store usefulness feedback without retaining the private checked content."""
    feedback = Feedback(
        helpful=payload.helpful,
        comment=f"kind={payload.kind}; result={payload.result_level}",
        contribution_consent=False,
    )
    db.add(feedback)
    db.commit()
    return {"id": str(feedback.id), "stored_content": False}


@router.get("/intelligence/users")
def list_users(db: Db, _: Analyst) -> list[dict]:
    users = db.scalars(
        select(User).options(selectinload(User.profile)).order_by(User.created_at.desc())
    ).all()
    return [serialize_user(user) for user in users]


@router.patch("/intelligence/users/{user_id}")
def update_user(user_id: UUID, payload: AdminUserUpdateRequest, db: Db, analyst: Analyst) -> dict:
    if analyst.profile.role != UserRole.super_admin:
        raise HTTPException(status_code=403, detail="Only a super administrator can change user access")
    user = db.scalar(
        select(User).options(selectinload(User.profile)).where(User.id == user_id)
    )
    if user is None:
        raise HTTPException(status_code=404, detail="User not found")
    if user.id == analyst.id and (
        payload.is_active is False
        or (payload.role is not None and payload.role != UserRole.super_admin)
    ):
        raise HTTPException(status_code=422, detail="You cannot remove your own super-administrator access")
    changes: dict[str, object] = {}
    if payload.role is not None and payload.role != user.profile.role:
        changes["role_from"] = user.profile.role.value
        changes["role_to"] = payload.role.value
        user.profile.role = payload.role
    if payload.is_active is not None and payload.is_active != user.is_active:
        changes["active_from"] = user.is_active
        changes["active_to"] = payload.is_active
        user.is_active = payload.is_active
    if changes:
        db.add(AuditLog(
            actor_id=analyst.id,
            action="user.access_updated",
            target_type="user",
            target_id=user.id,
            metadata_=changes,
        ))
        db.commit()
        db.refresh(user)
    return serialize_user(user)


@router.get("/intelligence/overview")
def intelligence_overview(db: Db, _: Analyst) -> dict:
    counts = {label.value: db.scalar(select(func.count()).select_from(DatasetItem).where(DatasetItem.label == label)) or 0 for label in DatasetLabel}
    report_categories = db.execute(
        select(
            FraudCategory.name,
            func.count(FraudReport.id),
            func.count(func.distinct(FraudReport.reporter_id)).filter(FraudReport.reporter_id.is_not(None)),
            func.count(DatasetItem.id).filter(DatasetItem.verification_status == ReviewStatus.verified),
        )
        .outerjoin(FraudReport, FraudReport.category_id == FraudCategory.id)
        .outerjoin(DatasetItem, DatasetItem.fraud_report_id == FraudReport.id)
        .group_by(FraudCategory.id, FraudCategory.name)
        .having(func.count(FraudReport.id) > 0)
        .order_by(func.count(FraudReport.id).desc(), FraudCategory.name)
    ).all()
    report_types = db.execute(
        select(FraudReport.entity_type, func.count(FraudReport.id))
        .group_by(FraudReport.entity_type)
        .order_by(func.count(FraudReport.id).desc())
    ).all()
    total_reports = db.scalar(select(func.count()).select_from(FraudReport)) or 0
    registered_reporters = db.scalar(
        select(func.count(func.distinct(FraudReport.reporter_id))).where(FraudReport.reporter_id.is_not(None))
    ) or 0
    training_ready = db.scalar(
        select(func.count()).select_from(DatasetItem).where(
            DatasetItem.verification_status == ReviewStatus.verified,
            DatasetItem.contribution_consent.is_(True),
            DatasetItem.label.in_((DatasetLabel.fraud, DatasetLabel.legitimate)),
        )
    ) or 0
    training_labels = {
        label.value: db.scalar(select(func.count()).select_from(DatasetItem).where(
            DatasetItem.verification_status == ReviewStatus.verified,
            DatasetItem.contribution_consent.is_(True),
            DatasetItem.label == label,
        )) or 0
        for label in (DatasetLabel.fraud, DatasetLabel.legitimate)
    }
    return {
        "total_items": db.scalar(select(func.count()).select_from(DatasetItem)) or 0,
        "pending_review": db.scalar(select(func.count()).select_from(ReviewQueueItem).where(ReviewQueueItem.status == ReviewStatus.pending)) or 0,
        "verified": db.scalar(select(func.count()).select_from(DatasetItem).where(DatasetItem.verification_status == ReviewStatus.verified)) or 0,
        "ghana_specific": db.scalar(select(func.count()).select_from(DatasetItem).where(DatasetItem.is_ghana_specific.is_(True))) or 0,
        "labels": counts,
        "reports": {
            "total": total_reports,
            "registered_reporters": registered_reporters,
            "anonymous_reports": db.scalar(select(func.count()).select_from(FraudReport).where(FraudReport.reporter_id.is_(None))) or 0,
            "training_ready": training_ready,
            "training_labels": training_labels,
            "by_type": [{"type": entity_type or "Unspecified", "count": count} for entity_type, count in report_types],
            "by_category": [
                {"category": name, "reports": report_count, "registered_reporters": reporter_count, "verified": verified_count}
                for name, report_count, reporter_count, verified_count in report_categories
            ],
        },
    }


@router.get("/intelligence/reports")
def intelligence_reports(db: Db, _: Analyst, limit: int = Query(100, ge=1, le=500)) -> list[dict]:
    rows = db.execute(
        select(
            FraudReport.id, FraudReport.created_at, FraudReport.entity_type, FraudReport.entity_value,
            FraudReport.description, FraudReport.status, FraudReport.reporter_id, FraudCategory.name,
            DatasetItem.verification_status, DatasetItem.channel,
        )
        .outerjoin(FraudCategory, FraudCategory.id == FraudReport.category_id)
        .outerjoin(DatasetItem, DatasetItem.fraud_report_id == FraudReport.id)
        .order_by(FraudReport.created_at.desc())
        .limit(limit)
    ).all()
    return [{
        "id": str(row.id), "created_at": row.created_at, "report_type": row.entity_type,
        "identifier": row.entity_value, "description": row.description, "status": row.status.value,
        "reporter": "registered" if row.reporter_id else "anonymous", "category": row.name or "Uncategorised",
        "review_status": row.verification_status.value if row.verification_status else "pending", "channel": row.channel,
    } for row in rows]


@router.get("/intelligence/review-queue")
def review_queue(db: Db, _: Analyst, queue_status: ReviewStatus = ReviewStatus.pending, limit: int = Query(50, ge=1, le=200)) -> list[dict]:
    entries = db.scalars(
        select(ReviewQueueItem).options(
            selectinload(ReviewQueueItem.dataset_item).selectinload(DatasetItem.source),
            selectinload(ReviewQueueItem.dataset_item).selectinload(DatasetItem.category),
        ).where(ReviewQueueItem.status == queue_status).order_by(ReviewQueueItem.created_at).limit(limit)
    ).all()
    return [{"queue_id": str(entry.id), "status": entry.status.value, **serialize_item(entry.dataset_item)} for entry in entries]


@router.post("/intelligence/review-queue/{queue_id}/decision")
def review_decision(queue_id: UUID, request: ReviewDecisionRequest, db: Db, analyst: Analyst) -> dict:
    entry = db.scalar(select(ReviewQueueItem).options(selectinload(ReviewQueueItem.dataset_item)).where(ReviewQueueItem.id == queue_id))
    if entry is None:
        raise HTTPException(status_code=404, detail="Review item not found")
    if request.decision == DatasetLabel.fraud and request.category_id is None:
        raise HTTPException(status_code=422, detail="A fraud category is required for a verified fraud item")
    entry.decision = request.decision
    entry.category_id = request.category_id
    entry.notes = request.notes
    entry.reviewed_by = analyst.id
    entry.reviewed_at = datetime.now(timezone.utc)
    entry.status = ReviewStatus.verified if request.approve else ReviewStatus.rejected
    entry.dataset_item.label = request.decision
    entry.dataset_item.category_id = request.category_id
    entry.dataset_item.verification_status = entry.status
    if entry.dataset_item.fraud_report_id:
        report = db.get(FraudReport, entry.dataset_item.fraud_report_id)
        if report is not None:
            if request.approve and request.decision == DatasetLabel.fraud:
                report.status = ReportStatus.verified_signal
                sync_verified_entity(db, entry.dataset_item, report)
            elif request.approve:
                report.status = ReportStatus.insufficient_evidence
            else:
                report.status = ReportStatus.rejected
    db.commit()
    return {"queue_id": str(entry.id), "status": entry.status.value, "dataset_item_id": str(entry.dataset_item_id)}


@router.get("/intelligence/dataset")
def list_dataset(
    db: Db, _: Analyst, label: DatasetLabel | None = None, source_type: DatasetSourceType | None = None,
    verified: bool | None = None, ghana_specific: bool | None = None, query: str | None = None,
    limit: int = Query(100, ge=1, le=500),
) -> list[dict]:
    statement = select(DatasetItem).options(selectinload(DatasetItem.source), selectinload(DatasetItem.category)).order_by(DatasetItem.created_at.desc())
    if label: statement = statement.where(DatasetItem.label == label)
    if source_type: statement = statement.join(DatasetSource).where(DatasetSource.source_type == source_type)
    if verified is not None: statement = statement.where(DatasetItem.verification_status == (ReviewStatus.verified if verified else ReviewStatus.pending))
    if ghana_specific is not None: statement = statement.where(DatasetItem.is_ghana_specific == ghana_specific)
    if query: statement = statement.where(DatasetItem.content.ilike(f"%{query}%"))
    return [serialize_item(item) for item in db.scalars(statement.limit(limit)).all()]


@router.get("/intelligence/training-datasets")
def training_datasets(db: Db, _: Analyst) -> list[dict]:
    rows = db.execute(
        select(DatasetVersion, func.count(DatasetVersionItem.dataset_item_id))
        .outerjoin(DatasetVersionItem, DatasetVersionItem.dataset_version_id == DatasetVersion.id)
        .group_by(DatasetVersion.id)
        .order_by(DatasetVersion.created_at.desc())
        .limit(25)
    ).all()
    return [{
        "id": str(version.id), "name": version.name, "version": version.version,
        "status": version.status, "notes": version.notes, "item_count": item_count,
        "created_at": version.created_at,
    } for version, item_count in rows]


@router.post("/intelligence/training-datasets", status_code=status.HTTP_201_CREATED)
def prepare_training_dataset(payload: TrainingDatasetRequest, db: Db, analyst: Analyst) -> dict:
    items = db.scalars(
        select(DatasetItem).where(
            DatasetItem.verification_status == ReviewStatus.verified,
            DatasetItem.contribution_consent.is_(True),
            DatasetItem.label.in_((DatasetLabel.fraud, DatasetLabel.legitimate)),
        ).order_by(DatasetItem.created_at)
    ).all()
    if not items:
        raise HTTPException(status_code=422, detail="Verify at least one fraud or legitimate item before preparing training data")
    version = DatasetVersion(
        name="guard-reviewed-intelligence",
        version=datetime.now(timezone.utc).strftime("v%Y%m%d-%H%M%S-%f"),
        status="ready",
        notes=payload.notes or "Human-reviewed, consented Guard intelligence",
        created_by=analyst.id,
    )
    db.add(version)
    db.flush()
    split_counts = {"train": 0, "validation": 0, "test": 0, "holdout": 0}
    label_counts = {"fraud": 0, "legitimate": 0}
    for item in items:
        bucket = int(item.content_hash[:8], 16) % 100
        split = "train" if bucket < 70 else "validation" if bucket < 85 else "test" if bucket < 95 else "holdout"
        split_counts[split] += 1
        label_counts[item.label.value] += 1
        db.add(DatasetVersionItem(dataset_version_id=version.id, dataset_item_id=item.id, split=split))
    db.add(AuditLog(
        actor_id=analyst.id, action="training_dataset.prepared", target_type="dataset_version",
        target_id=version.id, metadata_={"items": len(items), "splits": split_counts, "labels": label_counts},
    ))
    db.commit()
    return {
        "id": str(version.id), "name": version.name, "version": version.version, "status": version.status,
        "item_count": len(items), "splits": split_counts, "labels": label_counts,
        "trainable": label_counts["fraud"] > 0 and label_counts["legitimate"] > 0,
    }


@router.post("/intelligence/dataset", status_code=status.HTTP_201_CREATED)
def add_dataset_item(request: ManualDatasetRequest, db: Db, _: Analyst) -> dict:
    source = db.get(DatasetSource, request.source_id)
    if source is None:
        raise HTTPException(status_code=404, detail="Dataset source not found")
    try:
        item = create_candidate(
            db, content=request.content, source=source, label=request.label,
            category_id=request.category_id, channel=request.channel, country_code=request.country_code,
            is_ghana_specific=request.is_ghana_specific, consent=True,
            metadata={"notes": request.notes} if request.notes else {},
        )
        db.commit()
    except IntegrityError:
        db.rollback()
        raise HTTPException(status_code=409, detail="A matching dataset item already exists")
    return {"id": str(item.id), "status": "pending_review"}


def parse_upload(filename: str, payload: bytes) -> list[dict]:
    extension = filename.lower().rsplit(".", 1)[-1]
    if extension == "json":
        decoded = json.loads(payload.decode("utf-8-sig"))
        return decoded if isinstance(decoded, list) else decoded.get("items", [])
    if extension == "csv":
        return list(csv.DictReader(io.StringIO(payload.decode("utf-8-sig"))))
    if extension == "xlsx":
        from openpyxl import load_workbook
        sheet = load_workbook(io.BytesIO(payload), read_only=True, data_only=True).active
        rows = sheet.iter_rows(values_only=True)
        headers = [str(value).strip() if value is not None else "" for value in next(rows)]
        return [dict(zip(headers, row)) for row in rows]
    raise HTTPException(status_code=415, detail="Upload a CSV, XLSX or JSON file")


@router.post("/intelligence/dataset/import", status_code=status.HTTP_201_CREATED)
async def import_dataset(db: Db, _: Analyst, source_id: UUID, file: UploadFile = File(...)) -> dict:
    source = db.get(DatasetSource, source_id)
    if source is None:
        raise HTTPException(status_code=404, detail="Dataset source not found")
    payload = await file.read()
    if len(payload) > 10 * 1024 * 1024:
        raise HTTPException(status_code=413, detail="Dataset uploads are limited to 10 MB")
    try:
        rows = parse_upload(file.filename or "", payload)
    except (ValueError, json.JSONDecodeError, StopIteration) as error:
        raise HTTPException(status_code=422, detail=f"Could not parse dataset: {error}")
    imported, skipped, errors = 0, 0, []
    for index, row in enumerate(rows, start=2):
        content = str(row.get("content") or row.get("message") or "").strip()
        if not content:
            errors.append({"row": index, "error": "content/message is required"})
            continue
        raw_label = str(row.get("label") or row.get("fraud_status") or "uncertain").lower().strip()
        label = DatasetLabel(raw_label) if raw_label in {item.value for item in DatasetLabel} else DatasetLabel.uncertain
        category = category_for(db, str(row.get("category") or row.get("fraud_category") or "")) if (row.get("category") or row.get("fraud_category")) else None
        try:
            with db.begin_nested():
                create_candidate(
                    db, content=content, source=source, label=label,
                    category_id=category.id if category else None, channel=str(row.get("channel") or "") or None,
                    country_code=str(row.get("country") or row.get("country_code") or "").upper()[:2] or None,
                    is_ghana_specific=str(row.get("ghana_specific") or "").lower() in ("1", "true", "yes", "gh"),
                    consent=True, metadata={"import_file": file.filename, "import_row": index},
                )
            imported += 1
        except IntegrityError:
            skipped += 1
    db.commit()
    return {"imported": imported, "duplicates_skipped": skipped, "errors": errors[:100], "status": "pending_review"}


@router.get("/intelligence/categories")
def categories(db: Db, _: Analyst) -> list[dict]:
    return [{"id": str(item.id), "name": item.name, "slug": item.slug} for item in db.scalars(select(FraudCategory).order_by(FraudCategory.name)).all()]


@router.get("/intelligence/sources")
def sources(db: Db, _: Analyst) -> list[dict]:
    return [{"id": str(item.id), "name": item.name, "source_type": item.source_type.value, "description": item.description, "authoritative": item.is_authoritative} for item in db.scalars(select(DatasetSource).order_by(DatasetSource.name)).all()]


@router.get("/intelligence/brands")
def brands(db: Db, _: Analyst) -> list[dict]:
    return [{"id": str(item.id), "name": item.name, "country_code": item.country_code, "verified": item.verified, "official_domains": item.official_domains} for item in db.scalars(select(Brand).order_by(Brand.name)).all()]


@router.get("/intelligence/urls")
def urls(db: Db, _: Analyst, limit: int = Query(100, ge=1, le=500)) -> list[dict]:
    return [{"id": str(item.id), "url": item.url, "verdict": item.verdict.value, "risk_score": item.risk_score, "reports": item.report_count, "verified": item.verified} for item in db.scalars(select(UrlIntelligence).order_by(UrlIntelligence.risk_score.desc()).limit(limit)).all()]


@router.get("/intelligence/phone-numbers")
def phone_numbers(db: Db, _: Analyst, limit: int = Query(100, ge=1, le=500)) -> list[dict]:
    return [{"id": str(item.id), "sender_id": item.sender_id, "verdict": item.verdict.value, "risk_score": item.risk_score, "reports": item.report_count, "verified": item.verified} for item in db.scalars(select(PhoneIntelligence).order_by(PhoneIntelligence.risk_score.desc()).limit(limit)).all()]
